"""Builds the sample papers used for demos and tests: python samples/make_samples.py"""
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent

HONEST = """The Dust Bowl and Why It Still Matters

When I first heard about the Dust Bowl in high school I honestly thought it was just a few bad storms. It wasn't. For most of the 1930s, farmers across Oklahoma, Kansas and Texas watched their land literally blow away, and a lot of them lost everything.

My grandpa grew up near Guymon, and he used to say the dirt got into everything. Your food. Your bed. Even the inside of the clock on the wall. That stuck with me. It's one thing to read numbers in a book and a totally different thing to picture a kid sweeping dirt out of a kitchen every morning before school.

So what actually caused it? Part of it was weather, since a long drought hit right when prices for wheat collapsed. But a bigger part was people. Farmers had torn up the prairie grass to plant more wheat, and once the grass was gone there was nothing left to hold the soil down. As one encyclopedia puts it, "the exposed soil turned to dust and blew away in enormous clouds that sometimes blackened the sky for days" (Dust Bowl Encyclopedia).

I think the lesson is pretty simple, even if we don't always follow it. If you treat land like it will last forever, it won't. The government eventually figured that out and started paying farmers to plant trees and change how they plowed, and those changes helped a lot. We'd be smart to remember that today.
"""

HUMAN_PARA = ("I picked the Dust Bowl for this paper because my family is from western Oklahoma and I've "
              "heard stories about it my whole life. Honestly I didn't know much about the actual causes "
              "until I started reading for this class, so this was a good excuse to finally learn.")

AI_PARA = ("Furthermore, the Dust Bowl serves as a powerful testament to the intricate relationship between "
           "human activity and the natural environment. It is important to note that the crisis was not "
           "solely the result of drought; rather, it emerged from a multifaceted interplay of economic "
           "pressures, agricultural practices, and environmental vulnerabilities. Moreover, the widespread "
           "adoption of mechanized farming played a pivotal role in transforming the landscape of the Great "
           "Plains. Ultimately, this historical episode underscores the paramount importance of sustainable "
           "land management in fostering long-term ecological resilience.")

COPIED_PARA = ("During the previous decades, extensive deep plowing of the virgin topsoil of the Great Plains "
               "had displaced the native, deep-rooted grasses that normally trapped soil and moisture even "
               "during periods of drought and high winds.")

PARAPHRASED_PARA = ("Many of these migrants, commonly called Okies no matter which state they came from, "
                    "traveled west to California looking for farm work, only to discover that conditions "
                    "there were hardly better than those they had left behind.")

WEB_PARA = ("On April 14, 1935, a day remembered as Black Sunday, one of the worst storms swept across "
            "the southern plains and turned day into night.[3] The Soil Conservation Service was created the "
            "same year.[4][citation needed]")


def build_mixed(path: Path) -> None:
    d = docx.Document()
    normal = d.styles["Normal"].font
    normal.name, normal.size = "Times New Roman", Pt(12)
    d.add_heading("The Dust Bowl: Causes and Consequences", level=1)
    for text in (HUMAN_PARA, AI_PARA, COPIED_PARA, PARAPHRASED_PARA):
        d.add_paragraph(text)
    p = d.add_paragraph()
    run = p.add_run(WEB_PARA)  # pasted from a web page: different font, size and colour
    run.font.name, run.font.size = "Arial", Pt(10.5)
    run.font.color.rgb = RGBColor(0x20, 0x21, 0x22)
    d.add_paragraph("In the end, I learned a lot from this project and it made me look at my "
                    "grandpa's stories differently than I did before.")
    d.save(str(path))


if __name__ == "__main__":
    (HERE / "honest_paper.txt").write_text(HONEST, encoding="utf-8")
    build_mixed(HERE / "mixed_paper.docx")
    print("wrote honest_paper.txt and mixed_paper.docx")
