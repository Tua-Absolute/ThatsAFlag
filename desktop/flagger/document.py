"""Loading papers and splitting them into paragraphs, sentences and words.

Every piece of text keeps its character offsets into Document.text so the
checks can point at exactly what they flagged and the GUI can highlight it.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

PARA_SEP = "\n\n"


@dataclass
class Span:
    start: int
    end: int
    text: str


@dataclass
class FormatRun:
    """A stretch of .docx text with uniform formatting."""
    start: int
    end: int
    font: str | None
    size: float | None
    color: str | None
    highlighted: bool
    hyperlink: bool


@dataclass
class Paragraph(Span):
    style: str = "Normal"

    @property
    def is_heading(self) -> bool:
        s = self.style.lower()
        return s.startswith("heading") or s == "title" or s == "subtitle"


@dataclass
class Document:
    text: str
    paragraphs: list[Paragraph]
    format_runs: list[FormatRun] = field(default_factory=list)  # only for .docx
    source_name: str = "pasted text"


def from_text(text: str, source_name: str = "pasted text") -> Document:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = []
    # Paragraphs are separated by one or more blank lines; a single newline
    # stays inside a paragraph (it is itself a possible copy-paste artifact).
    for m in re.finditer(r"[^\n](?:.|\n(?!\s*\n))*", text):
        chunk = m.group(0)
        stripped = chunk.rstrip()
        if stripped.strip():
            paragraphs.append(Paragraph(m.start(), m.start() + len(stripped), stripped))
    return Document(text, paragraphs, source_name=source_name)


def load(path: str | Path) -> Document:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in (".txt", ".md", ""):
        return from_text(path.read_text(encoding="utf-8", errors="replace"), path.name)
    raise ValueError(f"Unsupported file type: {suffix} (use .txt, .md or .docx)")


def _load_docx(path: Path) -> Document:
    import docx  # python-docx
    from docx.text.hyperlink import Hyperlink
    from docx.text.run import Run

    d = docx.Document(str(path))
    default_font = d.styles["Normal"].font
    pieces: list[str] = []
    paragraphs: list[Paragraph] = []
    runs: list[FormatRun] = []
    pos = 0

    for p in d.paragraphs:
        p_start = pos
        p_text_parts = []
        style_font = p.style.font if p.style is not None else None
        for item in p.iter_inner_content():
            is_link = isinstance(item, Hyperlink)
            item_runs = item.runs if is_link else [item] if isinstance(item, Run) else []
            for r in item_runs:
                t = r.text
                if not t:
                    continue
                f = r.font
                font = f.name or (style_font.name if style_font else None) or default_font.name
                size_obj = f.size or (style_font.size if style_font else None) or default_font.size
                size = size_obj.pt if size_obj is not None else None
                color = None
                try:
                    if f.color is not None and f.color.type is not None and f.color.rgb is not None:
                        color = str(f.color.rgb)
                except (AttributeError, ValueError):
                    color = None
                rpr = r._element.rPr
                shaded = rpr is not None and rpr.find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd") is not None
                runs.append(FormatRun(pos, pos + len(t), font, size, color,
                                      bool(f.highlight_color) or shaded, is_link))
                p_text_parts.append(t)
                pos += len(t)
        p_text = "".join(p_text_parts)
        if p_text.strip():
            paragraphs.append(Paragraph(p_start, pos, p_text, p.style.name if p.style is not None else "Normal"))
        pieces.append(p_text)
        pieces.append(PARA_SEP)
        pos += len(PARA_SEP)

    return Document("".join(pieces), paragraphs, runs, path.name)


# --- sentences and words ---------------------------------------------------

_SENT_END = re.compile(r"(?<=[.!?])[\"'”’)\]]*\s+(?=[\"'“‘(\[]?[A-Z0-9])")
_WORD = re.compile(r"[A-Za-z0-9]+(?:['’][A-Za-z]+)?")


def sentences(span: Span) -> list[Span]:
    out = []
    start = 0
    text = span.text
    for m in _SENT_END.finditer(text):
        piece = text[start:m.start()].strip()
        if piece:
            s = text.index(piece, start)
            out.append(Span(span.start + s, span.start + s + len(piece), piece))
        start = m.end()
    piece = text[start:].strip()
    if piece:
        s = text.index(piece, start)
        out.append(Span(span.start + s, span.start + s + len(piece), piece))
    return out


def words(span: Span) -> list[Span]:
    return [Span(span.start + m.start(), span.start + m.end(), m.group(0))
            for m in _WORD.finditer(span.text)]


def norm(word: str) -> str:
    return word.lower().replace("’", "'")
