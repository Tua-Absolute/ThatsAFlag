"""Builds pasted_paper.docx: a paper carrying the .docx traces of pasting.
Run with: python make_pasted_paper.py  (needs python-docx)"""
import datetime as dt
import re
import shutil
import zipfile
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
PARAS = [
    "My family has farmed in western Oklahoma for four generations, and the Dust Bowl comes up at nearly every holiday dinner. I wanted to understand what actually happened instead of just repeating the stories I grew up with, so I chose it for this paper and started with the basics.",
    "Farmers on the southern plains plowed under millions of acres of native grassland during the wheat boom of the 1910s and 1920s. When drought arrived in 1931 the exposed topsoil had nothing to anchor it, and strong winds carried it hundreds of miles in rolling black clouds.",
    "By the middle of the decade, the government had begun paying farmers to adopt contour plowing, crop rotation and shelterbelts. These programs, run through the new Soil Conservation Service, are a big reason the region recovered when rain returned near 1940.",
    "Reading about all this made my grandpa's stories feel a lot more real to me. It's strange to think the same fields I drive past every summer were once blowing away, and I'm glad I finally took the time to learn the history behind them.",
]

d = docx.Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(11)
d.add_heading("What Caused the Dust Bowl", level=1)
for i, t in enumerate(PARAS):
    p = d.add_paragraph()
    r = p.add_run(t)
    r._element.set(qn("w:rsidR"), "00A1B2C3" if i != 1 else "00F00D01")
    if i == 1:
        # pasted from a web page: Word's "Normal (Web)" style and a British language tag
        p.style = d.styles.add_style("Normal (Web)", 1)
        rpr = r._element.get_or_add_rPr()
        lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "en-GB"); rpr.append(lang)
p = d.add_paragraph()
r = p.add_run("Every holiday my aunt still tells the story of the day the sky went black, and now I finally understand why it scared everybody so much back then.")
r._element.set(qn("w:rsidR"), "00A1B2C3")
cp = d.core_properties
cp.author, cp.last_modified_by, cp.revision = "Student", "Student", 1
cp.created = dt.datetime(2026, 10, 1, 21, 0, 0)
cp.modified = dt.datetime(2026, 10, 1, 21, 6, 0)
tmp = HERE / "_tmp.docx"
d.save(tmp)

words = sum(len(t.split()) for t in PARAS) * 2  # pretend it is a longer paper
app = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Template>Normal.dotm</Template><TotalTime>4</TotalTime><Words>{words}</Words><Application>Microsoft Office Word</Application><AppVersion>16.0000</AppVersion></Properties>"""
with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(HERE / "pasted_paper.docx", "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "docProps/app.xml":
            data = app.encode()
        zout.writestr(item, data)
tmp.unlink()
print("wrote pasted_paper.docx")
